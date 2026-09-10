import os
import sys
import json
from dotenv import load_dotenv
from google import genai
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

# Adjust path so we can import local modules from backend folder
sys.path.append(os.path.dirname(os.path.abspath(__file__)))
from extract import extract_text, extract_excel

# Load API key via python-dotenv
load_dotenv()
api_key = os.getenv("GEMINI_API_KEY")

if not api_key:
    print("Warning: GEMINI_API_KEY not found in .env file.")

def extract_fields(text):
    """
    Sends text to Gemini asking for ONLY valid JSON with specific fields.
    Parses and returns the JSON.
    """
    print("Extracting fields using Gemini...")
    client = genai.Client(api_key=api_key)
    
    prompt = f"""
    Analyze the following text and extract the key information.
    Return ONLY a valid JSON object with the following keys, strictly formatted.
    
    Required JSON schema:
    - "subsidiary_name" (string)
    - "report_year" (string)
    - "executive_summary" (string, multi-paragraph high-level overview)
    - "operational_and_production_metrics" (array of objects, each object containing: "category", "metric_name", "value", "unit", "yoy_change", "notes")
    - "financial_overview" (object containing: "revenue", "capital_expenditure", "operating_costs", "margins")
    - "exploration_and_development_projects" (array of objects, each object containing: "project_name", "location", "targets", "current_status")
    - "key_risks_and_challenges" (array of objects, each object containing: "risk_type", "description", "mitigation")
    - "strategic_outlook_and_targets" (string, next 1-3 year guidance)
    - "key_highlights" (list of strings, detailed bullet points with context)
    
    Do not include markdown blocks like ```json or ``` in your response, just the raw JSON object. Ensure all properties are present even if the value is null or empty.
    
    Text to analyze:
    {text[:30000]}
    """
    
    try:
        response = client.interactions.create(
            model='gemini-3.7-flash',
            input=prompt
        )
        response_text = response.output_text.strip()
        
        # Remove markdown formatting if Gemini included it
        if response_text.startswith("```json"):
            response_text = response_text[7:]
        if response_text.startswith("```"):
            response_text = response_text[3:]
        if response_text.endswith("```"):
            response_text = response_text[:-3]
            
        data = json.loads(response_text.strip())
        print("Successfully extracted and parsed JSON fields.")
        return data
    except Exception as e:
        print(f"Error during field extraction: {e}")
        # Return fallback structure on failure
        return {
            "subsidiary_name": "Unknown",
            "report_year": "Unknown",
            "executive_summary": "Data not available.",
            "operational_and_production_metrics": [],
            "financial_overview": {"revenue": "N/A", "capital_expenditure": "N/A", "operating_costs": "N/A", "margins": "N/A"},
            "exploration_and_development_projects": [],
            "key_risks_and_challenges": [],
            "strategic_outlook_and_targets": "Data not available.",
            "key_highlights": []
        }

def generate_report(file_path):
    """
    Extracts text/data via extract.py, calls extract_fields, and builds a Word document.
    """
    print(f"Generating report for: {file_path}")
    
    # 1. Extract text/data
    text = ""
    if file_path.lower().endswith('.pdf'):
        text = extract_text(file_path)
    elif file_path.lower().endswith(('.xlsx', '.csv')):
        text, _ = extract_excel(file_path)
    else:
        print(f"Unsupported file format for {file_path}. Use PDF, XLSX, or CSV.")
        return None
        
    if not text:
        print("Failed to extract any text from the file.")
        return None
        
    print(f"Extracted {len(text)} characters of text.")
    
    # 2. Call extract_fields
    fields = extract_fields(text)
    
    # 3. Build Word document
    print("Building Word document...")
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    
    # Title
    sub_name = fields.get('subsidiary_name', 'Unknown Subsidiary')
    report_year = fields.get('report_year', 'Unknown Year')
    
    title = doc.add_heading(f"Corporate Report: {sub_name}", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph(f"Report Year: {report_year}")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].bold = True
    subtitle.runs[0].font.size = Pt(14)
    
    # Removed page break so the first page isn't left blank
    doc.add_paragraph()
    
    # Executive Summary
    h1 = doc.add_heading("1. Executive Summary", level=1)
    h1.paragraph_format.page_break_before = False
    exec_sum = fields.get('executive_summary', 'N/A')
    p = doc.add_paragraph(str(exec_sum))
    
    doc.add_heading("Key Highlights", level=2)
    highlights = fields.get('key_highlights', [])
    if isinstance(highlights, list) and highlights:
        for h in highlights:
            doc.add_paragraph(str(h), style='List Bullet')
    else:
        doc.add_paragraph(str(highlights))
        
    doc.add_paragraph()
    
    # Operational Metrics (Table)
    doc.add_heading("2. Operational and Production Metrics", level=1)
    metrics = fields.get('operational_and_production_metrics', [])
    if isinstance(metrics, list) and metrics:
        table = doc.add_table(rows=1, cols=6)
        # Using a much cleaner, professional table style
        try:
            table.style = 'Light Shading Accent 1'
        except KeyError:
            table.style = 'Table Grid' # Fallback
            
        hdr_cells = table.rows[0].cells
        headers = ['Category', 'Metric', 'Value', 'Unit', 'YoY Change', 'Notes']
        for i, header in enumerate(headers):
            hdr_cells[i].text = header
            hdr_cells[i].paragraphs[0].runs[0].bold = True
            
        for metric in metrics:
            row_cells = table.add_row().cells
            row_cells[0].text = str(metric.get('category', ''))
            row_cells[1].text = str(metric.get('metric_name', ''))
            row_cells[2].text = str(metric.get('value', ''))
            row_cells[3].text = str(metric.get('unit', ''))
            row_cells[4].text = str(metric.get('yoy_change', ''))
            row_cells[5].text = str(metric.get('notes', ''))
    else:
        doc.add_paragraph("No operational metrics available.")
        
    doc.add_paragraph() # Replaced page break with a paragraph space
        
    # Financial Overview
    doc.add_heading("3. Financial Overview", level=1)
    finances = fields.get('financial_overview', {})
    if isinstance(finances, dict) and finances:
        for key, value in finances.items():
            clean_key = key.replace('_', ' ').title()
            p = doc.add_paragraph()
            p.add_run(f"{clean_key}: ").bold = True
            p.add_run(str(value))
    else:
        doc.add_paragraph("No financial data available.")
        
    # Exploration & Development Projects
    doc.add_heading("4. Exploration & Development Projects", level=1)
    projects = fields.get('exploration_and_development_projects', [])
    if isinstance(projects, list) and projects:
        for proj in projects:
            doc.add_heading(str(proj.get('project_name', 'Unnamed Project')), level=2)
            doc.add_paragraph(f"Location: {proj.get('location', 'N/A')}")
            doc.add_paragraph(f"Targets: {proj.get('targets', 'N/A')}")
            doc.add_paragraph(f"Status: {proj.get('current_status', 'N/A')}")
    else:
        doc.add_paragraph("No exploration projects available.")
        
    # Risks & Challenges
    doc.add_heading("5. Key Risks & Challenges", level=1)
    risks = fields.get('key_risks_and_challenges', [])
    if isinstance(risks, list) and risks:
        for risk in risks:
            doc.add_heading(str(risk.get('risk_type', 'General Risk')), level=2)
            doc.add_paragraph(f"Description: {risk.get('description', 'N/A')}")
            doc.add_paragraph(f"Mitigation: {risk.get('mitigation', 'N/A')}")
    else:
        doc.add_paragraph("No risk data available.")
        
    # Strategic Outlook
    doc.add_heading("6. Strategic Outlook & Targets", level=1)
    doc.add_paragraph(str(fields.get('strategic_outlook_and_targets', 'N/A')))
        
    # Footer citing the source file
    section = doc.sections[0]
    footer = section.footer
    footer_para = footer.paragraphs[0]
    footer_para.text = f"Source File: {os.path.basename(file_path)}"
    
    # Save to outputs/sample_report.docx
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    outputs_dir = os.path.join(base_dir, "outputs")
    os.makedirs(outputs_dir, exist_ok=True)
    
    output_path = os.path.join(outputs_dir, "sample_report.docx")
    try:
        doc.save(output_path)
        print(f"Report successfully saved to {output_path}")
        return output_path
    except Exception as e:
        print(f"Error saving report: {e}")
        return None

if __name__ == "__main__":
    # Find a sample file in the data folder to test
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    data_dir = os.path.join(base_dir, "data")
    
    if os.path.exists(data_dir):
        files = os.listdir(data_dir)
        if files:
            sample_file = os.path.join(data_dir, files[0])
            generate_report(sample_file)
        else:
            print("No files found in the data directory to test with.")
    else:
        print(f"Data directory not found at {data_dir}")
